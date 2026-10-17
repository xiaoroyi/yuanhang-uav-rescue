from pathlib import Path
import sys,json
ROOT=Path(__file__).resolve().parents[1];PROJECT=ROOT.parent;OUT=ROOT/'results';CACHE=ROOT/'cache'
sys.path.insert(0,str(PROJECT/'问题2/src'))
import env
sys.path.insert(0,str(ROOT/'deps311'))
from data_io import dump,table,sha
from prepare import load
from terrain import FORWARD,INVERSE
import numpy as np

def inputs():
    data=load()
    from openpyxl import load_workbook
    source=PROJECT.parent/'D题'
    f=next(p for p in source.rglob('通信链路参数.xlsx') if not p.name.startswith('~$'))
    w=load_workbook(f,data_only=True);vals={(r[0],r[3]):r[4] for r in w['数据'].values if len(r)>4 and isinstance(r[4],(int,float))};w.close()
    rr=next(p for p in source.rglob('中继无人机数据.xlsx') if not p.name.startswith('~$'))
    w=load_workbook(rr,data_only=True);rows=list(w['数据'].values);r=next(r for r in rows if r[0]=='R' and len(r)>4 and r[4]==23.5)
    keys=['m_empty','module','mass','speed','cruise_power','Euse','rho_percent','prep','link','turn','up','down','eta','down_eff','hover','radio','max_agl']
    relay=dict(zip(keys,r[2:19]));relay['K']=6;relay['M']=2;relay['Tfull']=1800
    assert any(r[0]=='R' and r[1:3]==(6,1800) for r in rows)
    radio={'frequency':vals['传播参数','f'],'system_loss':vals['传播参数','Lsys'],'obstacle':vals['传播参数','Lobs'],'sensitivity':vals['接收参数','Psens'],'margin':vals['接收参数','M']}
    pairs={'direct':('运输无人机','固定网关 G01'),'access':('运输无人机','中继接入端'),'backhaul':('中继回传端','固定网关 G01')}
    for name,(a,b) in pairs.items():
        ab=vals[a,'Pt']+vals[a,'G']+vals[b,'G']-radio['system_loss']-radio['sensitivity']-radio['margin']
        ba=vals[b,'Pt']+vals[b,'G']+vals[a,'G']-radio['system_loss']-radio['sensitivity']-radio['margin']
        radio[name]={'forward':ab,'reverse':ba,'limit':min(ab,ba)}
    for n in data['nodes'].values():n['xy']=list(FORWARD.transform(n['lon'],n['lat']))
    gateway=np.array(data['nodes']['O01']['xy']+[data['nodes']['O01']['z']+vals['固定网关 G01','hG']])
    data.update(relay=relay,radio=radio,gateway=gateway)
    data['audit']['inputs'].update({str(f):sha(f),str(rr):sha(rr)})
    return data

if __name__=='__main__':
    d=inputs();dump(OUT/'input_audit.json',d['audit']);dump(OUT/'config.json',dict(relay=d['relay'],radio=d['radio'],gateway=d['gateway'].tolist(),geometry='UTM49N projected DEM closed convex quadrilaterals; original heights; numerical cross-check of subdivided pixel edges',time_grid=.1,objective=['D','joint_return','total_energy','total_sorties']))
    print(json.dumps({'dem_shape':d['dem'].shape,'radio':d['radio'],'relay':d['relay'],'gateway':d['gateway'].tolist()},ensure_ascii=False))
