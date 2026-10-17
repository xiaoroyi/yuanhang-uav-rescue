import io
import json

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.app.main import create_app


def test_csv_and_xlsx_templates_roundtrip_and_keep_input_immutable(tmp_path):
    with TestClient(create_app(tmp_path / "runtime")) as client:
        s = client.post(
            "/api/v1/scenarios/examples", json={"example": "compact"}
        ).json()
        for kind, entity in [
            ("xlsx", "cargo_items"),
            ("csv", "cargo_items"),
            ("csv", "sites"),
            ("csv", "drone_types"),
            ("csv", "batteries"),
        ]:
            template = client.get(
                f"/api/v1/scenarios/{s['scenario_id']}/template?format={kind}&entity={entity}"
            )
            assert template.status_code == 200
            preview = client.post(
                f"/api/v1/scenarios/import-file?scenario_id={s['scenario_id']}&entity={entity}",
                files={"file": ("scene." + kind, template.content)},
            )
            assert preview.status_code == 200, preview.text
            imported = preview.json()["scenario"]
            assert imported["data_origin"] == "user_input"
            assert imported[entity] == s[entity]
            assert client.get("/api/v1/scenarios/" + s["scenario_id"]).json() == s
        workbook = load_workbook(
            io.BytesIO(
                client.get(f"/api/v1/scenarios/{s['scenario_id']}/template").content
            ),
            read_only=True,
        )
        assert len(workbook.sheetnames) == 7
        workbook.close()


def test_bad_row_identifies_line_and_never_creates_a_scene(tmp_path):
    with TestClient(create_app(tmp_path / "runtime")) as client:
        s = client.post(
            "/api/v1/scenarios/examples", json={"example": "compact"}
        ).json()
        response = client.post(
            f"/api/v1/scenarios/import-file?scenario_id={s['scenario_id']}&entity=demands",
            files={"file": ("bad.csv", b"demand_id,required_quantity\nD1,1.5\n")},
        )
        assert response.status_code == 422
        assert "2" in response.json()["error"]["message"]
        assert len(client.get("/api/v1/scenarios").json()) == 1
        oversized = client.post(
            "/api/v1/scenarios/import-file",
            files={"file": ("huge.json", b"x" * (12 * 1024 * 1024 + 1))},
        )
        assert oversized.status_code == 422


def test_nested_json_false_zero_and_missing_are_not_collapsed(tmp_path):
    with TestClient(create_app(tmp_path / "runtime")) as client:
        s = client.post(
            "/api/v1/scenarios/examples", json={"example": "compact"}
        ).json()
        s["batteries"][0]["soc"] = 0
        s["demands"][0]["expected_deadline_s"] = None
        s["policy"]["requires_terrain"] = False
        response = client.post(
            "/api/v1/scenarios/import-file",
            files={"file": ("new.json", json.dumps(s).encode())},
        )
        assert response.status_code == 200, response.text
        result = response.json()["scenario"]
        assert result["batteries"][0]["soc"] == 0
        assert result["demands"][0]["expected_deadline_s"] is None
        assert result["policy"]["requires_terrain"] is False
