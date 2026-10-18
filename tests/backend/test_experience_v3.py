"""Targeted V3 additions: template isolation, real workbook and spatial samples."""

import copy
import io
import json
from pathlib import Path

import openpyxl
import pytest
from fastapi.testclient import TestClient

from backend.app.main import create_app
from backend.app.schedule_export import schedule_workbook
from backend.app.radio_coverage import radio_coverage


@pytest.mark.parametrize("template", ["compact_joint", "contest_joint"])
def test_joint_template_creates_independent_inputs_and_is_idempotent(
    tmp_path, template
):
    with TestClient(create_app(tmp_path / "runtime")) as c:
        body = {
            "name": "Joint template",
            "template": template,
            "idempotency_key": template,
        }
        created = c.post("/api/v1/missions", json=body)
        assert created.status_code == 201, created.text
        m = created.json()
        assert c.post("/api/v1/missions", json=body).json() == m
        s = c.get("/api/v1/scenarios/" + m["current_scenario_id"]).json()
        assert s["policy"]["requires_communication"]
        assert len(s["communication"]["relay_drones"]) == 2
        assert len(s["communication"]["relay_energy"]) == 6
        assert len(s["cargo_items"]) == (32 if template == "compact_joint" else 80)
        other = c.post(
            "/api/v1/missions", json={**body, "idempotency_key": template + "-2"}
        ).json()
        assert other["mission_id"] != m["mission_id"]
        assert other["current_scenario_id"] != m["current_scenario_id"]
        assert c.get("/api/v1/scenarios/" + s["scenario_id"] + "/plans").json() == []


def joint_input():
    root = Path(__file__).resolve().parents[2]
    fixture = json.loads(
        (root / "evidence/solver/joint-compact.json").read_text(encoding="utf-8")
    )
    return fixture["snapshot"], fixture["outcome"]["plan"]


def test_workbook_preserves_numeric_seconds_sources_and_formula_like_names():
    snap, plan = joint_input()
    original = copy.deepcopy((snap, plan))
    snap["scenario"]["sites"][0]["name"] = "=1+1"
    source = copy.deepcopy(snap)
    book = openpyxl.load_workbook(
        io.BytesIO(schedule_workbook(plan, snap)), data_only=False
    )
    assert len(book.sheetnames) == 6
    assert book["运输架次"].max_row == len(plan["sorties"]) + 1
    assert (
        book["逐箱交付"].max_row
        == sum(len(r["cargo_ids"]) for r in plan["sorties"]) + 1
    )
    assert book["中继架次"].max_row == len(plan["relay_sorties"]) + 1
    assert book["通信保障"].max_row == len(plan["communication_report"]["windows"]) + 1
    assert book["运输架次"].freeze_panes == "A2"
    assert isinstance(book["运输架次"]["F2"].value, (float, int))
    assert not any(c.data_type == "f" for sheet in book for row in sheet for c in row)
    assert any(plan["plan_id"] == c.value for row in book["方案说明"] for c in row)
    assert snap == source and plan == original[1]


def test_spatial_coverage_is_saved_source_bound_and_rejects_foreign_relay():
    snap, plan = joint_input()
    original = copy.deepcopy((snap, plan))
    rid = plan["relay_sorties"][0]["relay_sortie_id"]
    coverage = radio_coverage(snap, plan, 120, rid, 9)
    assert len(coverage["points"]) == 81
    assert coverage["snapshot_id"] == snap["snapshot_id"]
    assert coverage["terrain_sha256"] == snap["scenario"]["terrain"]["sha256"]
    assert coverage["relay_backhaul_available"] is True
    assert (
        coverage["relay_service_start_s"] == plan["relay_sorties"][0]["service_start_s"]
    )
    assert set(p["source"] for p in coverage["points"]) <= {
        "gateway",
        "relay",
        "uncovered",
        "unknown",
    }
    assert any(p["source"] == "relay" for p in coverage["points"])
    with pytest.raises(KeyError):
        radio_coverage(snap, plan, 120, "foreign-relay", 9)
    assert (snap, plan) == original
