"""Generate a report from completed saved evidence, without running any search.

Use the bundled document Python runtime (python-docx), not the app environment.
"""
from pathlib import Path
import csv
import json
import math
import hashlib
from datetime import datetime
from zoneinfo import ZoneInfo

ROOT=Path(__file__).resolve().parents[2]
EVIDENCE=ROOT/'evidence/contest-baselines-v1'
BASE=EVIDENCE/'matrix_180s_seed42'
DEST=ROOT/'docs/current-project/contest-baselines-v1'

def read(path):
    return json.loads(path.read_text(encoding='utf-8'))

def label(q,m):
    names={'greedy':'Greedy / EDF','aco':'ACO','alns':'ALNS','cp_sat':'CP-SAT','cp_sat_warm':'CP-SAT 贪心初解'}
    name=names[m]
    return name+' + CP排程' if q=='Q3' and m not in ('cp_sat','cp_sat_warm') else name

def f(value,n=3):
    return '—' if value is None else f'{value:.{n}f}'

def groups(records,refs):
    output=[]
    for q in ['Q2','Q3']:
        for obj,title,suffix in [('time','时间优先','fast'),('energy','能耗优先','economy')]:
            h=refs[q+'_'+suffix]
            hs=dict(completion_s=float(h['C']),energy_kwh=float(h['E']),
                    transport_sorties=int(h['N_transport'] or h['N']),relay_sorties=int(h['N_relay'] or 0))
            selected=[r for r in records if r['question']==q and r['objective']==obj]
            selected.sort(key=lambda r:['greedy','aco','alns','cp_sat','cp_sat_warm'].index(r['method']))
            rows=[]
            for r in selected:
                v=r['metrics'] or {}
                method_name=label(q,r['method'])
                if obj=='energy' and r['method']=='cp_sat':method_name+=' 冷启动'
                rows.append([method_name,f(v.get('completion_s')),f(v.get('energy_kwh')),
                             '—' if not v else f"{v['transport_sorties']} + {v['relay_sorties']}",
                             f(r['search_wall_s'],2),'通过' if v else '未得完整解'])
            rows.append(['论文历史参考',f(hs['completion_s']),f(hs['energy_kwh']),f"{hs['transport_sorties']} + {hs['relay_sorties']}",'未匹配','历史记录'])
            feasible=[r for r in selected if r['metrics']]
            metric='completion_s' if obj=='time' else 'energy_kwh'
            if feasible:
                best=min(feasible,key=lambda r:r['metrics'][metric])
                value=best['metrics'][metric];gain=(value-hs[metric])/value*100
                text=(f"在本次已得可行基线中，{label(q,best['method'])}的{'完成时间' if obj=='time' else '总能耗'}最低，"
                      f"为{value:.3f}{'秒' if obj=='time' else 'kWh'}。论文参考值为{hs[metric]:.3f}，"
                      f"相对该基线{'降低' if gain>=0 else '增加'}{abs(gain):.2f}%。"
                      '这一差异同时受搜索预算、候选空间及算法适配影响，不能归因于算法本身。')
            else:
                text='本组限时基线未获得通过完整核验的方案；不据此判定原题不可行。'
            output.append(dict(question=q,objective=obj,title=q+' '+title,rows=rows,analysis=text,
                               historical=hs,best_baseline=None if not feasible else dict(method=best['method'],value=value,reference_reduction_percent=gain)))
    return output


def main():
    records=read(BASE/'results.json');complete=read(BASE/'completion.json');checks=read(BASE/'final_checks.json')
    assert complete['complete'] and checks['all_checks_passed'] and len(records)==16
    warm=read(EVIDENCE/'cp_energy_warm_180s/results.json')
    warm_completion=read(EVIDENCE/'cp_energy_warm_180s/completion.json')
    assert len(warm)==2 and warm_completion['complete']
    for r in warm:
        r['method']='cp_sat_warm'
    records=records+warm
    refs={r['scenario']:r for r in read(EVIDENCE/'historical_reference.json')}
    gs=groups(records,refs)
    cross=[]
    for q in ['Q2','Q3']:
        feasible=[r for r in records if r['question']==q and r['metrics']]
        for metric,obj,suffix in [('completion_s','时间','fast'),('energy_kwh','能耗','economy')]:
            winner=min(feasible,key=lambda r:r['metrics'][metric])
            hist=float(refs[q+'_'+suffix]['C' if metric=='completion_s' else 'E'])
            value=winner['metrics'][metric]
            cross.append(dict(question=q,target=obj,method=label(q,winner['method']),
                run_objective=winner['objective'],baseline_value=value,historical_value=hist,
                historical_reduction_percent=(value-hist)/value*100))
    DEST.mkdir(exist_ok=True,parents=True)
    valid=sum(r['metrics'] is not None for r in records)
    intro=(f'本报告比较原题第二问和第三问的运输及通信联合调度，分别以完成时间和总能耗为目标。'
           f'本次首轮运行16个配置，另补充2个CP-SAT贪心初解配置，合计18次运行，其中{valid}次获得通过完整核验的方案。'
           '我们的方法按用户要求直接引用最终论文的时间优先和节能备选结果，没有重新搜索。'
           '本次基线适合补充为初步对照，尚不足以证明方法在同预算、多随机种子条件下普遍领先。')
    protocol=[
        '数据与任务：1个基地、15个服务区、80个不可拆货箱、8架异构运输机、14组运输电池；Q3另含2架中继机与6个能源组件。',
        '完整约束：载质量与体积、逐段剩余载荷能耗、DEM净空、返航SOC不低于20%、同型共享电池和两阶段充电、实体资源占用、全部31箱硬截止。',
        '计时口径：同次停靠全部货箱的交接结束时刻为本站全部箱件的交付时刻。全部80箱整站期望延期为零；没有混用早期逐箱计时。',
        '时间目标：Q2最晚运输返航；Q3运输与中继全部架次的最晚返航。能耗目标：Q2运输总能耗；Q3运输加中继总能耗，单位kWh。架次数仅为辅助指标。',
        '能耗组没有额外设置完成时间上限；论文Q3节能备选是在99分钟以内得到的历史方案，因此保留其原始条件说明。',
        '公共输入：850条运输候选，最多3个服务区停靠；21个由服务节点和4公里网格最高DEM点自动生成的300米离地中继候选，至多4个中继架次槽。没有读取论文的路线、分组、时刻或驻点作初解。',
        '计算设置：固定种子42；迭代方法每配置180秒墙钟搜索预算，包含构造和模型建立；停止检测、结果解码与写盘可能略超预算，实际耗时逐行列出，最终独立核验另计。四个独立进程并发，每个CP求解器1个线程；设备和其它后台负载未严格控制。',
        '贪心不迭代，完成一次构造即停止；Q3首次公共CP排程最多20秒。ACO与ALNS后续每个候选的CP排程最多12秒，所有子调用累计计入本配置180秒。',
        '两个节能CP冷启动返回UNKNOWN后，各补跑一次180秒，使用同一普通时间贪心提供可行初解；Q3初解由最多20秒的时间CP排程建立。没有使用论文初解，原冷启动失败记录保留。',
    ]
    algorithms=[
        ('Greedy / EDF','每轮选择最早到期货箱，以完成时间/装载量或能耗/装载量评分，从公共候选集中选可容纳的架次；显式分配机体和电池。构造失败如实保留。'),
        ('ACO','在EDF焦点的前10个可行候选中，按信息素与启发式的概率抽样；信息素指数1、启发式指数3、蒸发率0.1，以当前最好路线沉积信息素。属于本题的蚁群式适配，非官方ACS逐行复现。'),
        ('ALNS','使用随机移除、单位箱件高能耗移除、空间相关移除三种破坏算子；破坏比例15%至40%，以带随机性的可行重插修复。每25轮更新算子权重，更新率0.2；采用降温接受。'),
        ('CP-SAT','在同一有限路线池中进行精确覆盖选择和资源累计排程，采用0.1秒保守时间网格和整数能耗。时间组保留公共贪心初解；能耗组同时列出冷启动与补充的时间贪心初解版本。限时结果只说明有限模型，不代表原题全局最优。'),
    ]
    q3_note=('Q3前三类使用共享CP排程模块，决定给定货箱分组的运输时序、中继驻点、服务窗口及能源分配，'
             '所以表中明确标注“+ CP排程”。CP-SAT行同时优化完整公共池的路线选择和联合排程。'
             '连续通信使用原论文的投影DEM闭像元几何模型；检查所有通信开区间和边界事件，并以GEOS独立射线抽检。')
    verification=(f'全部有数值的结果均从保存的方案重建80箱唯一交付、载荷、物理能耗、返航电量、机体与电池占用及充电；'
                  '额外单独核验整站零延期。Q3还通过连续通信及中继机体、能源组件检查，通信缺口和未确定区间均为零。'
                  '负例测试验证了重复箱件、缺箱、篡改能耗、时限违反及删除中继服务均被识别。'
                  '有结果不等于最优；“未得完整解”不等于原题无解。')
    limitations=[
        '论文参考不是同预算实验。论文Q3最终一轮时间阶段约403.886秒、能耗阶段约721.165秒，另有前序搜索和候选构造；本次单配置预算180秒。不得写成“三分钟内优于所有传统方法”。',
        '每配置只有一个随机种子，没有均值、标准差、置信区间或显著性检验，也不是16个独立问题实例。',
        '公共池和驻点集有限，方法适配和候选质量会影响排名。原论文使用更丰富的搜索过程，故差异包含候选空间差异。',
        '论文参考是历史研究方案，不是本轮V4网页引擎的运行输出；这些数值不能用于宣称当前软件在180秒内达到同样成绩。',
        'Q3的ACO和ALNS为混合实现，不能宣传为纯ACO、纯ALNS或原作者官方实现。CP-SAT亦是有限池模型基线，论文方法本身也使用CP-SAT，不能表述为“击败CP-SAT求解器”。',
        '贪心节能方向的构造失败保留在分母。比较只对完整可行方案计算，缺失指标不填0，也不为改善结果放松时限、载荷或通信要求。',
        'DeepACO未训练或运行，PSO和WOA未实现。本轮优先可解释、可复核的传统方法，不以方法数量替代有效对照。',
    ]
    budget=(f"公共预处理实测{read(EVIDENCE/'preparation.json')['preprocessing_wall_s']:.2f}秒；"
            f"正式16配置矩阵实测墙钟{complete['experiment_wall_s']/60:.2f}分钟。"
            f"两项补充初解实验墙钟{warm_completion['wall_s']/60:.2f}分钟。"
            f"首轮各配置累计搜索墙钟{complete['sum_search_wall_s']/60:.2f}分钟，含核验的配置耗时累计{complete['sum_total_run_wall_s']/60:.2f}分钟。"
            '并发后累计耗时与实际等待不同。12个迭代配置的预算合计36分钟；贪心提早停止。'
            '两项补充实验的搜索预算另计6分钟，总搜索预算42分钟。'
            '若首轮迭代配置都延长至300秒，预算合计60分钟；同等4进程条件预计首轮矩阵约15至20分钟，未在本轮执行。'
            '若使用3个种子、每次180秒，则迭代预算108分钟，4进程约27分钟起，另加预处理与核验。')
    references=[
        '原题和最终提交论文及答案：用户指定的华为杯2026 D题目录；已核对论文、答案、5张输入表和DEM共8份原件的SHA-256与仓库存档完全一致。',
        '论文正式结果：物理第43页的Q2两目标图及第61页表5.14；完整精度由正式发布accepted_metrics.csv保留，Q2时间优先另与答案Excel逐项核对。',
        'Dorigo M, Gambardella L M. Ant Colony System: A Cooperative Learning Approach to the Traveling Salesman Problem. IEEE Transactions on Evolutionary Computation, 1997. https://iridia.ulb.ac.be/~mdorigo/ACO/publications.html',
        'Ropke S, Pisinger D. An Adaptive Large Neighborhood Search Heuristic for the Pickup and Delivery Problem with Time Windows. Transportation Science, 2006, 40(4):455–472. https://doi.org/10.1287/trsc.1050.0135',
        'Google OR-Tools. CP-SAT Solver and Setting Solver Limits. https://developers.google.com/optimization/cp/cp_solver ; https://developers.google.com/optimization/cp/cp_tasks',
    ]
    paragraphs=[g['analysis'] for g in gs]
    competition=('在镇龙乡15个服务区、80个不可拆货箱的标准算例上，我们设置了时间优先与能耗优先两类目标，'
                 '并构造贪心、蚁群式搜索、自适应大邻域搜索和有限池CP-SAT对照。全部可行结果均采用整站交接完成计时，'
                 '满足硬截止、容量、返航安全电量和资源周转约束，联合问题另核验全过程通信。'
                 '论文历史方案展示了时效与能耗的可行取舍；本次短时基线为其提供初步量化参照。'
                 '由于历史方案与新基线的总搜索预算及候选空间不同，当前证据用于展示已得方案质量，'
                 '不作同预算普遍优势或原问题全局最优的结论。')
    md=['# 援航原题第二问第三问限时基线对比实验报告','',
        '日期：2026年10月19日　任务：EXP-C01　主矩阵：matrix_180s_seed42','',intro,'','## 实验口径','']
    md.extend('- '+x for x in protocol)
    md+=['','## 基线实现','']
    for name,desc in algorithms:md += [f'**{name}**：{desc}','']
    md += [q3_note,'','## 结果对比','']
    header=['方法','全部返航 s','总能耗 kWh','运输 + 中继架次','搜索耗时 s','状态']
    for g in gs:
        md += ['### '+g['title'],'','| '+' | '.join(header)+' |','|'+'|'.join(['---']*len(header))+'|']
        md += ['| '+' | '.join(row)+' |' for row in g['rows']]
        md += ['',g['analysis'],'']
    md += ['### 汇总全部已得可行方案','',
        '有限搜索中，设置能耗目标不保证得到比时间组更低能耗的解。为避免忽略较强对照，下表在同一道题的全部已运行配置（含1个补充CP初解配置）中分别选取最低完成时间与最低能耗。每个单元格对应一份真实完整方案，未拼接不同方案的时间与能耗。该汇总使用了多次运行的结果，不是额外一个180秒算法。','',
        '| 题目与指标 | 最强已得基线来源 | 基线值 | 论文参考 | 参考值相对变化 |',
        '|---|---|---|---|---|']
    for row in cross:
        gain=row['historical_reduction_percent']
        md += [f"| {row['question']} {row['target']} | {row['method']}（{row['run_objective']}运行） | {row['baseline_value']:.3f} | {row['historical_value']:.3f} | {'降低' if gain>=0 else '增加'}{abs(gain):.2f}% |"]
    md += ['## 核验与可复现性','',verification,'',
        '完整方案、求解状态、轨迹、核验明细和输入指纹保存在 `evidence/contest-baselines-v1/matrix_180s_seed42/`，补充初解实验在 `cp_energy_warm_180s/`。',
        '公共物理与连续通信数学函数由只读数模档案选择性复制，替换旧路径与导入后运行，没有执行归档bootstrap/env。来源与原文件哈希见 `source_manifest.json`。',
        '', '## 时间预算','',budget,'','## 适用边界','']
    md += ['- '+x for x in limitations]
    md += ['','## 可用于项目文档的表述','',competition,'','## 复现命令','',
        '```powershell', '.venv/Scripts/python.exe scripts/contest_baselines/prepare_sources.py',
        '.venv/Scripts/python.exe scripts/contest_baselines/benchmark.py prepare',
        '.venv/Scripts/python.exe scripts/contest_baselines/verify_inputs.py',
        '# matrix.py使用新输出目录；重新运行前更改其NAME，保留旧证据',
        '.venv/Scripts/python.exe scripts/contest_baselines/matrix.py',
        '.venv/Scripts/python.exe scripts/contest_baselines/check_experiment.py --final --name matrix_180s_seed42',
        '```','', '配置以protocol.json为准。墙钟限时和后台负载会改变实际迭代数，复现不承诺逐位相同结果；保存方案可直接重新核验。',
        '', '## 资料来源','']
    md += ['- '+x for x in references]
    report=DEST/'对比实验报告.md';report.write_text('\n'.join(md)+'\n',encoding='utf-8')
    # A five-section companion document, using the same saved numbers and prose.
    from docx import Document
    from docx.shared import Inches,Pt,RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc=Document();section=doc.sections[0]
    section.page_width=Inches(8.5);section.page_height=Inches(11)
    section.top_margin=section.bottom_margin=Inches(.65)
    section.left_margin=section.right_margin=Inches(.65)
    normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(10.5)
    normal.element.rPr.rFonts.set(qn('w:eastAsia'),'宋体')
    normal.paragraph_format.space_after=Pt(6);normal.paragraph_format.line_spacing=1.12
    for name in ['Title','Heading 1','Heading 2']:
        st=doc.styles[name];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
        st.element.rPr.rFonts.set(qn('w:eastAsia'),'黑体')
    doc.styles['Title'].font.size=Pt(19)
    doc.styles['Heading 1'].font.size=Pt(14)
    doc.styles['Heading 2'].font.size=Pt(12)
    # Bundled Word styles contain a blue Title border; remove inherited rules.
    for style in doc.styles:
        for border in list(style.element.xpath('.//w:pBdr')):
            border.getparent().remove(border)
    doc.core_properties.title='援航原题第二问第三问限时基线对比实验报告'
    doc.core_properties.author='援航项目组'
    def para(text,style=None):return doc.add_paragraph(text,style)
    def heading(text):doc.add_heading(text,1)
    def word_table(g):
        doc.add_heading(g['title'],2)
        tab=doc.add_table(rows=1,cols=6);tab.style='Table Grid';tab.alignment=WD_TABLE_ALIGNMENT.CENTER
        tab.autofit=False
        widths=[1.76,1.10,1.08,1.07,1.02,1.17]
        for cell,w,text in zip(tab.rows[0].cells,widths,['方法','全部返航\n秒','总能耗\nkWh','运输＋中继\n架次','搜索耗时\n秒','结果状态']):
            cell.width=Inches(w);cell.text=text
            shading=OxmlElement('w:shd');shading.set(qn('w:fill'),'E8EEF4');cell._tc.get_or_add_tcPr().append(shading)
        repeat=OxmlElement('w:tblHeader');tab.rows[0]._tr.get_or_add_trPr().append(repeat)
        for row in g['rows']:
            cells=tab.add_row().cells
            for cell,w,text in zip(cells,widths,row):cell.width=Inches(w);cell.text=text
        for row in tab.rows:
            no_split=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no_split)
            for i,cell in enumerate(row.cells):
                cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for p in cell.paragraphs:
                    p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(4)
                    p.alignment=WD_ALIGN_PARAGRAPH.LEFT if i==0 else WD_ALIGN_PARAGRAPH.CENTER
                    for run in p.runs:run.font.size=Pt(9.5)
        para(g['analysis'])
    doc.add_heading('援航原题第二问第三问\n限时基线对比实验报告',0)
    para('2026年10月19日  |  EXP-C01  |  180秒固定种子初步对照')
    para(intro)
    heading('实验口径')
    for text in protocol:para(text)
    doc.add_page_break();heading('第二问运输调度结果')
    for g in gs[:2]:word_table(g)
    para('两组结果都要求全部80箱整站按期交付。论文参考为已提交历史方案，本轮没有重新搜索；其总计算预算与本次基线未匹配。')
    heading('核验范围');para(verification)
    doc.add_page_break();heading('第三问运输与中继联合结果')
    for g in gs[2:]:word_table(g)
    para(q3_note)
    para('表中的搜索耗时包含本配置构造与CP模型建立；最终连续通信、资源和物理独立核验另计。未得完整解不填0，也不等于原题无解。')
    doc.add_page_break();heading('基线实现与证据边界')
    for name,desc in algorithms:para(name+'：'+desc)
    for text in limitations:para(text)
    doc.add_page_break();heading('全部已得方案与时间预算')
    para('为避免忽略较强对照，在同一道题的全部已运行配置（含补充CP初解配置）中分别汇总最低完成时间与最低能耗。各指标来自真实完整方案，没有拼接方案；此项汇总使用了多次运行，不是额外一个180秒算法。')
    for row in cross:
        gain=row['historical_reduction_percent']
        para(f"{row['question']} {row['target']}：最强已得基线为{row['method']}的{row['run_objective']}运行，{row['baseline_value']:.3f}；论文参考{row['historical_value']:.3f}，相对{'降低' if gain>=0 else '增加'}{abs(gain):.2f}%。")
    heading('时间预算');para(budget)
    heading('项目文档建议表述');para(competition)
    heading('复现与来源')
    para('完整代码、配置、来源哈希、保存方案及核验记录见同名Markdown报告和 evidence/contest-baselines-v1。论文与答案原件未改写。')
    for text in references[2:]:
        p=para(text)
        for run in p.runs:run.font.size=Pt(8)
    footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    run=footer.add_run('援航实验报告  ·  ');run.font.size=Pt(9)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    word=DEST/'援航限时基线对比实验报告.docx';doc.save(word)
    dump_data=dict(groups=gs,cross_objective_best=cross,validated_configurations=valid,total_configurations=18,
        report_sha256=hashlib.sha256(report.read_bytes()).hexdigest(),docx_sha256=hashlib.sha256(word.read_bytes()).hexdigest())
    (BASE/'report_summary.json').write_text(json.dumps(dump_data,ensure_ascii=False,indent=2),encoding='utf-8')
    with (DEST/'完整结果对比.csv').open('w',encoding='utf-8-sig',newline='') as file:
        keys=['question','objective','method','status','budget_s','search_wall_s','completion_s','energy_kwh','transport_sorties','relay_sorties','minimum_hard_slack_s']
        writer=csv.DictWriter(file,fieldnames=keys);writer.writeheader()
        for r in records:
            writer.writerow({k:r.get(k,(r['metrics'] or {}).get(k)) for k in keys})
    print(str(report));print(str(word));print(json.dumps(gs,ensure_ascii=False,indent=2))

if __name__=='__main__':main()
