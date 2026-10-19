"""Isolated read-only contest inputs; all caches belong to this experiment."""
from pathlib import Path
import csv
import json
from functools import lru_cache
import numpy as np
from openpyxl import load_workbook
from scipy.io import loadmat
from pyproj import Transformer

PROJECT = Path(__file__).resolve().parents[2]
EVIDENCE = PROJECT/'evidence/contest-baselines-v1'
OUT = EVIDENCE/'inputs'
CACHE = EVIDENCE/'cache'
FORWARD = Transformer.from_crs(4326, 32649, always_xy=True)
INVERSE = Transformer.from_crs(32649, 4326, always_xy=True)

def dump(path, data):
    path = Path(path)
    path.parent.mkdir(exist_ok=True, parents=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2, allow_nan=False), encoding='utf-8')

def table(path, rows):
    rows = list(rows)
    if not rows:
        return
    with Path(path).open('w', encoding='utf-8-sig', newline='') as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)

@lru_cache(maxsize=1)
def inputs():
    data = json.loads((OUT/'normalized_inputs.json').read_text(encoding='utf-8'))
    mat = loadmat(next((PROJECT/'references/modeling/data/terrain').glob('*.mat')))
    data.update(dem=mat['dem'], lon=mat['longitude'].ravel(), lat=mat['latitude'].ravel())
    for n in data['nodes'].values():
        n['xy'] = list(FORWARD.transform(n['lon'], n['lat']))
    tables = PROJECT/'references/modeling/data/tables'
    w = load_workbook(tables/'通信链路参数.xlsx', read_only=True, data_only=True)
    vals = {(r[0], r[3]):r[4] for r in w['数据'].values if len(r)>4 and isinstance(r[4], (int, float))}
    w.close()
    w = load_workbook(tables/'中继无人机数据.xlsx', read_only=True, data_only=True)
    row = next(r for r in w['数据'].values if r[0]=='R' and len(r)>4 and r[4]==23.5)
    w.close()
    keys = ['m_empty','module','mass','speed','cruise_power','Euse','rho_percent','prep','link','turn','up','down','eta','down_eff','hover','radio','max_agl']
    relay = dict(zip(keys, row[2:19]))
    relay.update(K=6, M=2, Tfull=1800)
    radio = {'frequency':vals['传播参数','f'], 'system_loss':vals['传播参数','Lsys'], 'obstacle':vals['传播参数','Lobs'], 'sensitivity':vals['接收参数','Psens'], 'margin':vals['接收参数','M']}
    for kind, a, b in [('direct','运输无人机','固定网关 G01'), ('access','运输无人机','中继接入端'), ('backhaul','中继回传端','固定网关 G01')]:
        ab = vals[a,'Pt']+vals[a,'G']+vals[b,'G']-radio['system_loss']-radio['sensitivity']-radio['margin']
        ba = vals[b,'Pt']+vals[b,'G']+vals[a,'G']-radio['system_loss']-radio['sensitivity']-radio['margin']
        radio[kind] = dict(forward=ab, reverse=ba, limit=min(ab, ba))
    base = data['nodes']['O01']
    data.update(relay=relay, radio=radio, gateway=np.array(base['xy']+[base['z']+vals['固定网关 G01','hG']]))
    return data
