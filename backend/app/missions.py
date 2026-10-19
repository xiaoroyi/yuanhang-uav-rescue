"""Mission aggregate and necessary-condition reports, additive to the V1 API."""

from __future__ import annotations
import copy
from datetime import datetime, timezone
import uuid
import csv
import io
import json
import zipfile

from fastapi import File, UploadFile, Query
from fastapi.responses import Response

from sqlalchemy import select
from pydantic import ValidationError

from backend.app.generated.domain import Scenario
from backend.app.jobs import now
from backend.app.operation_models import (
    MissionCreate,
    MissionUpdate,
    MissionDraft,
    MissionDraftUpdate,
    MissionView,
    PreflightRequest,
    PreflightReport,
    StageAnalysis,
    DraftImportPreview,
)
from backend.app.storage import (
    MissionRow,
    OperationRecordRow,
    ScenarioRow,
    PlanRow,
    JobRow,
    ExecutionRow,
    SnapshotRow,
)
from shared.contracts import content_hash, validate_scenario, ContractError
from solver.adapters.contest import compact_scenario, contest_scenario
from solver.physics import parameters, leg, InputNotReady
from solver.stages import stage_analysis
from solver.weather import type_checks as weather_type_checks
from backend.app.imports import GROUPS, parse_rows
from backend.app.generated.domain import (
    Site,
    Demand,
    CargoItem,
    DroneType,
    Drone,
    Battery,
)


def identity(prefix):
    return prefix + "-" + uuid.uuid4().hex[:16].upper()


def draft_from(s):
    return MissionDraft.model_validate(
        {k: v for k, v in s.items() if k in MissionDraft.model_fields}
    ).model_dump(mode="json", exclude_unset=True)


def require(session, model, key, error):
    row = session.get(model, key)
    if row is None:
        raise error(404, "REFERENCE_INVALID", "未找到指定任务或执行记录。")
    return row


def revision(content, expected, error):
    if content["revision"] != expected:
        raise error(409, "STATE_VERSION_CONFLICT", "任务或执行已更新，请刷新后重试。")


def mission_view(session, row):
    view = copy.deepcopy(row.content)
    scenario = session.get(ScenarioRow, row.scenario_id) if row.scenario_id else None
    if scenario:
        view["draft"] = draft_from(scenario.content)
        active_plan = (
            session.get(PlanRow, scenario.active_plan_id)
            if scenario.active_plan_id
            else None
        )
        view["active_plan_id"] = (
            active_plan.id
            if active_plan and active_plan.version == scenario.version
            else None
        )
        view["site_count"] = len(scenario.content["sites"])
        view["cargo_count"] = len(scenario.content["cargo_items"])
        view["timezone"] = scenario.content["timezone"]
        view["updated_at"] = max(view["updated_at"], scenario.updated_at)
    else:
        view["site_count"] = len(view["draft"]["sites"])
        view["cargo_count"] = len(view["draft"]["cargo_items"])
    archived = view["lifecycle"] == "archived"
    if not archived:
        ex = (
            session.get(ExecutionRow, view["active_execution_id"])
            if view["active_execution_id"]
            else None
        )
        if ex:
            view["lifecycle"] = {
                "ready": "ready",
                "running": "executing",
                "paused": "paused",
                "completed": "completed",
            }[ex.content["status"]]
            view["next_action"] = (
                "review" if ex.content["status"] == "completed" else "execute"
            )
        elif scenario:
            jobs = list(
                session.scalars(select(JobRow).where(JobRow.scenario_id == scenario.id))
            )
            plans = list(
                session.scalars(
                    select(PlanRow).where(PlanRow.scenario_id == scenario.id)
                )
            )
            if view["active_plan_id"]:
                view.update(lifecycle="ready", next_action="execute")
            elif plans or any(j.state in ("queued", "running") for j in jobs):
                view.update(lifecycle="planning", next_action="plans")
    return view


def register_missions(app, database, error, save_new, snapshot_for):
    group_types = dict(
        zip(GROUPS, [Site, Demand, CargoItem, DroneType, Drone, Battery])
    )

    @app.post(
        "/api/v1/missions/{mission_id}/input-preview", response_model=DraftImportPreview
    )
    async def input_preview(
        mission_id: str,
        entity: str = Query("cargo_items"),
        file: UploadFile = File(...),
    ):
        if entity not in GROUPS:
            raise error(422, "SCHEMA_INVALID", "请选择合法实体表。")
        data = await file.read(12 * 1024 * 1024 + 1)
        if len(data) > 12 * 1024 * 1024:
            raise error(422, "SCHEMA_INVALID", "文件超出12MiB导入限制。")
        with database.session() as session:
            m = mission_view(session, require(session, MissionRow, mission_id, error))
        draft = copy.deepcopy(m["draft"])
        try:
            suffix = (file.filename or "").lower().rsplit(".", 1)[-1]
            if suffix == "csv":
                draft[entity] = parse_rows(
                    list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))
                )
            elif suffix == "json":
                value = json.loads(data.decode("utf-8-sig"))
                if isinstance(value, list):
                    draft[entity] = value
                elif isinstance(value, dict):
                    value = value.get("scenario", value.get("draft", value))
                    draft.update(
                        {
                            k: v
                            for k, v in value.items()
                            if k in MissionDraft.model_fields
                        }
                    )
                else:
                    raise ValueError("请输入任务输入对象或当前实体数组。")
            elif suffix == "xlsx":
                from openpyxl import load_workbook

                with zipfile.ZipFile(io.BytesIO(data)) as archive:
                    entries = archive.infolist()
                    if (
                        len(entries) > 2000
                        or sum(x.file_size for x in entries) > 50 * 1024 * 1024
                        or any(".." in x.filename.split("/") for x in entries)
                    ):
                        raise ValueError("工作簿解压规模或路径不合法。")
                wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
                try:
                    metadata = (
                        "mission_draft"
                        if "mission_draft" in wb.sheetnames
                        else "scenario"
                    )
                    if metadata in wb.sheetnames:
                        rows = parse_rows([list(r) for r in wb[metadata].values])
                        if len(rows) != 1:
                            raise ValueError("元数据工作表应恰好一行。")
                        draft.update(
                            {
                                k: v
                                for k, v in rows[0].items()
                                if k in MissionDraft.model_fields
                            }
                        )
                    for group in GROUPS:
                        if group in wb.sheetnames:
                            draft[group] = parse_rows(
                                [list(r) for r in wb[group].values]
                            )
                finally:
                    wb.close()
            else:
                raise ValueError("仅支持UTF-8 CSV、JSON和XLSX。")
            normalized = MissionDraft.model_validate(draft).model_dump(
                mode="json", exclude_unset=True
            )
        except (
            ValueError,
            TypeError,
            KeyError,
            ValidationError,
            zipfile.BadZipFile,
        ) as exc:
            raise error(
                422, "IMPORT_INVALID", "导入预览失败：" + str(exc)[:1800]
            ) from exc
        return {
            "draft": normalized,
            "row_count": sum(len(normalized[g]) for g in GROUPS),
            "notices": [
                "预览尚未保存；发布输入时校验所有实体引用，求解时独立检查物理约束。"
            ],
        }

    @app.get("/api/v1/missions/{mission_id}/input-template")
    def input_template(
        mission_id: str, entity: str = "cargo_items", format: str = "xlsx"
    ):
        if entity not in GROUPS or format not in ("csv", "xlsx"):
            raise error(422, "SCHEMA_INVALID", "不支持的模板类型。")
        with database.session() as session:
            draft = mission_view(
                session, require(session, MissionRow, mission_id, error)
            )["draft"]

        def cells(row, keys):
            return [
                (
                    json.dumps(row.get(k), ensure_ascii=False)
                    if isinstance(row.get(k), (list, dict))
                    else row.get(k)
                )
                for k in keys
            ]

        if format == "csv":
            stream = io.StringIO()
            writer = csv.writer(stream)
            keys = list(group_types[entity].model_fields)
            writer.writerow(keys)
            for row in draft[entity]:
                writer.writerow(cells(row, keys))
            return Response(
                ("\ufeff" + stream.getvalue()).encode("utf-8"),
                media_type="text/csv",
                headers={
                    "Content-Disposition": f'attachment; filename="mission-{entity}.csv"'
                },
            )
        from openpyxl import Workbook

        wb = Workbook()
        wb.remove(wb.active)
        metadata = {k: v for k, v in draft.items() if k not in GROUPS}
        for group, rows, keys in [
            ("mission_draft", [metadata], list(metadata)),
            *[(g, draft[g], list(group_types[g].model_fields)) for g in GROUPS],
        ]:
            ws = wb.create_sheet(group)
            ws.append(keys)
            for row in rows:
                ws.append(cells(row, keys))
            ws.freeze_panes = "A2"
            for column in ws.columns:
                ws.column_dimensions[column[0].column_letter].width = 24
        stream = io.BytesIO()
        wb.save(stream)
        wb.close()
        return Response(
            stream.getvalue(),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={
                "Content-Disposition": 'attachment; filename="yuanhang-mission.xlsx"'
            },
        )

    @app.get("/api/v1/missions", response_model=list[MissionView])
    def list_missions():
        with database.session() as session:
            return sorted(
                [mission_view(session, r) for r in session.scalars(select(MissionRow))],
                key=lambda m: m["updated_at"],
                reverse=True,
            )

    @app.post("/api/v1/missions", response_model=MissionView, status_code=201)
    def create_mission(body: MissionCreate):
        digest = content_hash(
            body.model_dump(
                mode="json", exclude={"case_id"} if body.case_id is None else set()
            )
        )
        if body.case_id and body.scenario_id:
            raise error(422, "SCHEMA_INVALID", "案例与已有场景不能同时作为来源。")
        with database.session(immediate=True) as session:
            prior = session.get(
                OperationRecordRow, ("mission.create", body.idempotency_key)
            )
            if prior:
                if prior.request_hash != digest:
                    raise error(
                        409, "IDEMPOTENCY_CONFLICT", "相同创建键已用于不同任务。"
                    )
                return mission_view(
                    session, require(session, MissionRow, prior.owner_id, error)
                )
            s = None
            if body.scenario_id:
                existing = session.scalar(
                    select(MissionRow).where(MissionRow.scenario_id == body.scenario_id)
                )
                if existing:
                    session.add(
                        OperationRecordRow(
                            namespace="mission.create",
                            id=body.idempotency_key,
                            owner_id=existing.id,
                            request_hash=digest,
                            content={"mission_id": existing.id},
                        )
                    )
                    return mission_view(session, existing)
                s = copy.deepcopy(
                    require(session, ScenarioRow, body.scenario_id, error).content
                )
            elif body.template != "blank" or body.case_id:
                from backend.app.storage import VersionRow

                s = (
                    compact_scenario()
                    if body.template in ("compact", "compact_joint")
                    else contest_scenario()
                )
                if body.case_id:
                    from backend.app.case_library import case_scenario

                    try:
                        s = case_scenario(body.case_id)
                    except ValueError as exc:
                        raise error(422, "SCHEMA_INVALID", str(exc)) from exc
                if body.template in ("compact_joint", "contest_joint"):
                    from solver.adapters.contest import contest_communication

                    s["communication"] = contest_communication(s)
                    s["policy"]["requires_communication"] = True
                s.update(
                    scenario_id=identity("SCN"),
                    state_version=1,
                    name=body.name,
                    mode=body.mode,
                )
                s = Scenario.model_validate(s).model_dump(
                    mode="json", exclude_unset=True
                )
                validate_scenario(s)
                session.add(
                    ScenarioRow(
                        id=s["scenario_id"],
                        name=s["name"],
                        version=1,
                        content=s,
                        active_plan_id=None,
                        updated_at=now(),
                    )
                )
                session.add(
                    VersionRow(
                        scenario_id=s["scenario_id"],
                        version=1,
                        content=s,
                        content_hash=content_hash(s),
                    )
                )
            template = compact_scenario()
            draft = (
                draft_from(s)
                if s
                else draft_from(
                    {
                        **template,
                        "base": None,
                        "sites": [],
                        "demands": [],
                        "cargo_items": [],
                        "drones": [],
                        "batteries": [],
                        "terrain": {
                            "status": "not_evaluated",
                            "asset_id": None,
                            "sha256": None,
                            "horizontal_crs": None,
                            "vertical_datum": None,
                        },
                        "policy": {**template["policy"], "requires_terrain": False},
                        "start_time": now(),
                    }
                )
            )
            mid = identity("MIS")
            value = {
                "mission_id": mid,
                **({"source_case_id": body.case_id} if body.case_id else {}),
                "name": body.name,
                "mode": s["mode"] if s else body.mode,
                "timezone": draft["timezone"],
                "lifecycle": "draft",
                "revision": 1,
                "current_scenario_id": s["scenario_id"] if s else None,
                "active_plan_id": None,
                "active_execution_id": None,
                "draft": draft,
                "updated_at": now(),
                "site_count": len(draft["sites"]),
                "cargo_count": len(draft["cargo_items"]),
                "next_action": "prepare",
            }
            row = MissionRow(
                id=mid, scenario_id=value["current_scenario_id"], content=value
            )
            session.add(row)
            session.add(
                OperationRecordRow(
                    namespace="mission.create",
                    id=body.idempotency_key,
                    owner_id=mid,
                    request_hash=digest,
                    content={
                        "mission_id": mid,
                        **({"source_case_id": body.case_id} if body.case_id else {}),
                    },
                )
            )
            return value

    @app.get("/api/v1/missions/{mission_id}", response_model=MissionView)
    def get_mission(mission_id: str):
        with database.session() as session:
            return mission_view(
                session, require(session, MissionRow, mission_id, error)
            )

    @app.patch("/api/v1/missions/{mission_id}", response_model=MissionView)
    def update_mission(mission_id: str, body: MissionUpdate):
        with database.session(immediate=True) as session:
            row = require(session, MissionRow, mission_id, error)
            revision(row.content, body.expected_revision, error)
            value = copy.deepcopy(row.content)
            if body.name is not None:
                value["name"] = body.name
            if body.archived is not None:
                if body.archived and value["active_execution_id"]:
                    ex = require(
                        session, ExecutionRow, value["active_execution_id"], error
                    )
                    if ex.content["status"] == "running":
                        raise error(409, "EXECUTION_ACTIVE", "请先暂停演练再归档。")
                value["lifecycle"] = "archived" if body.archived else "draft"
            value.update(revision=value["revision"] + 1, updated_at=now())
            row.content = value
            return mission_view(session, row)

    @app.put("/api/v1/missions/{mission_id}/draft", response_model=MissionView)
    def update_draft(mission_id: str, body: MissionDraftUpdate):
        with database.session(immediate=True) as session:
            row = require(session, MissionRow, mission_id, error)
            revision(row.content, body.expected_revision, error)
            if row.scenario_id:
                raise error(
                    409,
                    "DRAFT_PUBLISHED",
                    "已发布输入请通过原版本接口编辑，保留并发保护。",
                )
            value = copy.deepcopy(row.content)
            value.update(
                draft=body.draft.model_dump(mode="json", exclude_unset=True),
                revision=value["revision"] + 1,
                updated_at=now(),
            )
            if body.publish:
                candidate = {
                    "schema_version": "0.1.0-draft",
                    "scenario_id": identity("SCN"),
                    "name": value["name"],
                    "mode": value["mode"],
                    "state_version": 1,
                    "data_origin": "user_input",
                    "coordinate_system": "EPSG:4326",
                    **value["draft"],
                }
                try:
                    s = Scenario.model_validate(candidate).model_dump(
                        mode="json", exclude_unset=True
                    )
                    validate_scenario(s)
                except (ValidationError, ContractError) as exc:
                    raise error(
                        422,
                        "INPUT_INCOMPLETE",
                        "请补齐基地、需求点、需求和机型；并检查引用。",
                        details={"reason": str(exc)[:2000]},
                    ) from exc
                # This transaction stores the first Scenario as well as its immutable version.
                from backend.app.storage import VersionRow

                session.add(
                    ScenarioRow(
                        id=s["scenario_id"],
                        name=s["name"],
                        version=1,
                        content=s,
                        active_plan_id=None,
                        updated_at=now(),
                    )
                )
                session.add(
                    VersionRow(
                        scenario_id=s["scenario_id"],
                        version=1,
                        content=s,
                        content_hash=content_hash(s),
                    )
                )
                row.scenario_id = s["scenario_id"]
                value["current_scenario_id"] = s["scenario_id"]
            row.content = value
            return mission_view(session, row)

    @app.post(
        "/api/v1/missions/{mission_id}/clone",
        response_model=MissionView,
        status_code=201,
    )
    def clone_mission(mission_id: str, body: MissionCreate):
        with database.session() as session:
            original = mission_view(
                session, require(session, MissionRow, mission_id, error)
            )
        created = create_mission(
            MissionCreate(
                name=body.name, mode="simulation", idempotency_key=body.idempotency_key
            )
        )
        if created["current_scenario_id"]:
            return created
        return update_draft(
            created["mission_id"],
            MissionDraftUpdate(
                expected_revision=created["revision"],
                draft=MissionDraft.model_validate(original["draft"]),
                publish=bool(original["current_scenario_id"]),
            ),
        )

    @app.post(
        "/api/v1/missions/{mission_id}/preflight",
        response_model=PreflightReport,
        status_code=201,
    )
    def preflight(mission_id: str, body: PreflightRequest):
        with database.session(immediate=True) as session:
            row = require(session, MissionRow, mission_id, error)
            revision(row.content, body.expected_revision, error)
            m = mission_view(session, row)
            scenario = (
                session.get(ScenarioRow, row.scenario_id) if row.scenario_id else None
            )
            s = scenario.content if scenario else m["draft"]
            checks = []

            def add(
                code,
                group,
                status,
                message,
                action,
                ids=(),
                actual=None,
                required=None,
                unit=None,
            ):
                checks.append(
                    {
                        "code": code,
                        "group": group,
                        "status": status,
                        "message": message,
                        "affected_objects": list(ids),
                        "action": action,
                        "source_ref": "mission_input:" + content_hash(s),
                        "field_path": action,
                        "actual": actual,
                        "required": required,
                        "unit": unit,
                    }
                )

            add(
                "INPUT_PUBLISHED",
                "model",
                "passed" if scenario else "blocked",
                "输入版本已保存" if scenario else "输入仍为草稿，请补齐并发布",
                "checks",
            )
            base = s.get("base")
            add(
                "BASE_AVAILABLE",
                "map",
                "passed" if base and base["available"] else "blocked",
                "基地可用" if base and base["available"] else "请设置可用起降基地",
                "region",
            )
            if not s["sites"] or not s["demands"]:
                add(
                    "DEMAND_EMPTY",
                    "inventory",
                    "blocked",
                    "请添加需求点与物资需求",
                    "demand",
                )
            for d in s["demands"]:
                count = sum(
                    c["demand_id"] == d["demand_id"]
                    and c["status"] in ("available", "reserved")
                    for c in s["cargo_items"]
                )
                add(
                    "INVENTORY:" + d["demand_id"],
                    "inventory",
                    "passed" if count >= d["required_quantity"] else "blocked",
                    f"{d['demand_id']}：库存 {count} / 需求 {d['required_quantity']} 箱",
                    "demand",
                    [d["demand_id"], d["site_id"]],
                    count,
                    d["required_quantity"],
                    "箱",
                )
            available = [d for d in s["drones"] if d["status"] == "available"]
            add(
                "FLEET_AVAILABLE",
                "fleet",
                "passed" if available else "blocked",
                f"可用机体 {len(available)} 架",
                "fleet",
                [d["drone_id"] for d in available],
            )
            for d in available:
                compatible = [
                    b
                    for b in s["batteries"]
                    if b["compatible_type_id"] == d["type_id"]
                    and b["status"] in ("available", "charging")
                    and b["soh"] > 0
                    and base
                    and b["location_id"] == base["id"]
                ]
                add(
                    "BATTERY:" + d["drone_id"],
                    "fleet",
                    "passed" if compatible else "blocked",
                    f"{d['drone_id']} 有 {len(compatible)} 块可用兼容电池",
                    "fleet",
                    [d["drone_id"]],
                )
            for t in s["drone_types"]:
                try:
                    parameters(t)
                    add(
                        "ENERGY_MODEL:" + t["type_id"],
                        "model",
                        "passed",
                        f"{t['type_id']} 能耗/充电模型可读取",
                        "fleet",
                        [t["type_id"]],
                    )
                except InputNotReady as exc:
                    add(
                        "ENERGY_MODEL:" + t["type_id"],
                        "model",
                        "blocked",
                        str(exc),
                        "fleet",
                        [t["type_id"]],
                    )
            if base:
                missing = [
                    n["id"]
                    for n in [base, *s["sites"]]
                    if n["ground_elevation_m"] is None
                ]
                add(
                    "ELEVATION",
                    "map",
                    "blocked" if missing else "passed",
                    "缺少点位高程" if missing else "点位高程完整",
                    "region",
                    missing,
                )
            required = s["policy"]["requires_terrain"]
            add(
                "TERRAIN",
                "map",
                (
                    "passed"
                    if s["terrain"]["status"] == "available"
                    else "blocked" if required else "unknown"
                ),
                (
                    "地形资产已绑定，覆盖将在航段计算复核"
                    if s["terrain"]["status"] == "available"
                    else "尚未绑定地形资产"
                ),
                "region",
            )
            if scenario and base and not missing:
                for site in s["sites"]:
                    try:
                        if not s["drone_types"]:
                            break
                        leg(s, s["drone_types"][0], base["id"], site["id"])
                    except InputNotReady as exc:
                        add(
                            "MAP_COVERAGE:" + site["id"],
                            "map",
                            "blocked",
                            str(exc),
                            "region",
                            [site["id"]],
                        )
            if s.get("weather"):
                weather = s["weather"]
                add(
                    "WEATHER_SOURCE",
                    "weather",
                    "passed",
                    f"人工统一情景：{weather['source_ref']}；窗口{weather['valid_from_s']:g}—{weather['valid_until_s']:g}s，仅检查作业限值",
                    "fleet",
                )
                rules = weather_type_checks(s)
                any_allowed = any(r["allowed"] for r in rules)
                for r in rules:
                    add(
                        "WEATHER_TYPE:" + r["type_id"],
                        "weather",
                        (
                            "passed"
                            if r["allowed"]
                            else "warning" if any_allowed else "blocked"
                        ),
                        r["type_id"]
                        + (
                            " 满足所配置气象限值"
                            if r["allowed"]
                            else " 将不参与排程：" + "；".join(r["issues"])
                        ),
                        "fleet",
                        [r["type_id"]],
                    )
                if s["policy"]["requires_communication"]:
                    add(
                        "WEATHER_COMBINATION",
                        "weather",
                        "blocked",
                        "气象窗口与联合通信组合尚未支持，不能忽略其中任一约束。",
                        "fleet",
                    )
            else:
                add(
                    "WEATHER_NOT_CONFIGURED",
                    "weather",
                    "unknown",
                    "未绑定气象；当前为显式无风题设参数演练",
                    "fleet",
                )
            comm = s["policy"]["requires_communication"]
            add(
                "COMMUNICATION",
                "communication",
                (
                    "warning"
                    if comm and s.get("communication")
                    else "blocked" if comm else "not_applicable"
                ),
                (
                    "通信配置已启用，完整联合核验在求解时完成"
                    if comm and s.get("communication")
                    else "必需通信配置缺失" if comm else "本次未启用通信约束"
                ),
                "fleet",
            )
            snap = None
            if body.check_kind != "planning_readiness":
                p = session.get(PlanRow, body.plan_id) if body.plan_id else None
                valid = bool(
                    p
                    and scenario
                    and p.scenario_id == scenario.id
                    and p.version == scenario.version
                    and p.confirmed_at
                    and scenario.active_plan_id == p.id
                )
                add(
                    "PLAN_CONFIRMED",
                    "execution",
                    "passed" if valid else "blocked",
                    "已确认当前版本方案" if valid else "请先确认当前输入版本的方案",
                    "plans",
                )
                if body.check_kind == "real_dispatch":
                    add(
                        "REAL_DISPATCH_DISABLED",
                        "execution",
                        "blocked",
                        "未配置真实设备适配、权限与现场许可",
                        "settings",
                    )
            report = {
                "report_id": identity("CHK"),
                "mission_id": mission_id,
                "scenario_id": row.scenario_id,
                "state_version": scenario.version if scenario else None,
                "snapshot_id": snap,
                "check_kind": body.check_kind,
                "checks": checks,
                "blocking_count": sum(c["status"] == "blocked" for c in checks),
                "warning_count": sum(
                    c["status"] in ("warning", "unknown") for c in checks
                ),
                "validity_fingerprint": content_hash(s),
                "checked_at": now(),
            }
            session.add(
                OperationRecordRow(
                    namespace="preflight",
                    id=report["report_id"],
                    owner_id=mission_id,
                    request_hash=content_hash(report),
                    content=report,
                )
            )
            return report

    @app.get("/api/v1/preflight-reports/{report_id}", response_model=PreflightReport)
    def get_report(report_id: str):
        with database.session() as session:
            return require(
                session, OperationRecordRow, ("preflight", report_id), error
            ).content

    @app.get(
        "/api/v1/missions/{mission_id}/preflight-reports",
        response_model=list[PreflightReport],
    )
    def list_reports(mission_id: str):
        with database.session() as session:
            require(session, MissionRow, mission_id, error)
            return sorted(
                [
                    r.content
                    for r in session.scalars(
                        select(OperationRecordRow).where(
                            OperationRecordRow.namespace == "preflight",
                            OperationRecordRow.owner_id == mission_id,
                        )
                    )
                ],
                key=lambda r: r["checked_at"],
                reverse=True,
            )

    @app.get("/api/v1/plans/{plan_id}/stages", response_model=StageAnalysis)
    def stages(plan_id: str):
        with database.session() as session:
            p = require(session, PlanRow, plan_id, error)
            snap = require(session, SnapshotRow, p.snapshot_id, error).content
            try:
                return stage_analysis(snap, p.content)
            except InputNotReady as exc:
                raise error(422, "UNSUPPORTED_SCOPE", str(exc)) from exc
