"""Readable answer-style schedules from the exact saved Plan and its Snapshot.

This is an application export, not a copy of the historical contest workbook.
"""

from __future__ import annotations

from io import BytesIO
import json
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter


def schedule_workbook(plan: dict, snapshot: dict) -> bytes:
    from solver.stages import stage_analysis

    scenario = snapshot["scenario"]
    phases = (
        stage_analysis(snapshot, plan)["trajectory_segments"]
        if plan.get("communication_report")
        else []
    )
    wb = Workbook()
    wb.remove(wb.active)
    sites = {x["id"]: x["name"] for x in [scenario["base"], *scenario["sites"]]}
    drones = {x["drone_id"]: x["type_id"] for x in scenario["drones"]}
    cargo = {x["cargo_id"]: x for x in scenario["cargo_items"]}
    demands = {x["demand_id"]: x for x in scenario["demands"]}
    types = {"MED": "医疗物资", "WAT": "饮用水", "FOD": "食品", "HYG": "卫生用品"}

    def label(site):
        return f"{site} · {sites.get(site, site)}"

    def sheet(name, headers, rows):
        ws = wb.create_sheet(name)
        ws.append(headers)
        for row in rows:
            ws.append(row)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        ws.sheet_view.showGridLines = False
        ws.row_dimensions[1].height = 32
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.print_title_rows = "1:1"
        for col, heading in enumerate(headers, 1):
            width = 28 if any(s in heading for s in ["编号", "顺序", "来源"]) else 20
            if any(s in heading for s in ["列表", "说明", "内容"]):
                width = 60
            ws.column_dimensions[get_column_letter(col)].width = width
        for row in ws:
            for cell in row:
                # Preserve text as text, including user IDs beginning with '='.
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                cell.border = Border(bottom=Side(style="hair", color="E1E8EE"))
                if cell.row == 1:
                    cell.fill = PatternFill("solid", fgColor="123C45")
                    cell.font = Font(
                        name="Microsoft YaHei", size=10, bold=True, color="FFFFFF"
                    )
                else:
                    cell.fill = PatternFill(
                        "solid", fgColor="F2F7F7" if cell.row % 2 == 0 else "FFFFFF"
                    )
                    cell.font = Font(name="Microsoft YaHei", size=10, color="243846")
                    if isinstance(cell.value, float):
                        cell.number_format = "0.000"
            if row[0].row > 1:
                ws.row_dimensions[row[0].row].height = 32
        return ws

    metrics = plan.get("joint_metrics") or plan["metrics"]
    sheet(
        "方案说明",
        ["项目", "内容"],
        [
            ["任务名称", scenario["name"]],
            ["方案编号", plan["plan_id"]],
            ["输入快照编号", plan["snapshot_id"]],
            ["输入版本", plan["scenario_state_version"]],
            [
                "调度偏好",
                {
                    "time": "时间优先",
                    "energy": "能耗优先",
                    "balanced": "时间与能耗折中",
                }.get(plan["objective"], plan["objective"]),
            ],
            ["计算来源", plan["data_origin"]],
            ["独立检查", plan["validation"]["status"]],
            ["任务起点", scenario["start_time"]],
            ["时区", scenario["timezone"]],
            [
                "时间口径",
                "各表时刻为任务启动后的相对秒数；准备、起飞、交付、返航、资源再次可用是不同事件。",
            ],
            [
                "交付口径",
                "以保存方案中的交付完成时刻为准；一次停靠的整站交接不等于灾区所有批次均完成。",
            ],
            [
                "能源口径",
                "能耗单位kWh；联合总能耗包含运输与通信中继；电量百分比仅用于比较安全储备。",
            ],
            ["执行口径", "本文件为计划安排，不是实际飞行或模拟已交付的事实记录。"],
            [
                "全部返航（秒）",
                metrics.get("joint_makespan_s", metrics.get("makespan_s")),
            ],
            [
                "总能耗（kWh）",
                metrics.get("joint_energy_kwh", metrics.get("energy_kwh")),
            ],
            ["目标规则", plan.get("objective_policy", "deadline_first_v1")],
            ["物理输入", json.dumps(scenario["policy"], ensure_ascii=False)],
        ],
    )
    ordered = sorted(
        plan["sorties"],
        key=lambda r: (r["prepare_start_s"], r["drone_id"], r["sortie_id"]),
    )
    sheet(
        "运输架次",
        [
            "序号",
            "架次编号",
            "无人机编号",
            "机型",
            "电池编号",
            "准备开始（秒）",
            "起飞时刻（秒）",
            "访问救援点顺序",
            "返回调度中心（秒）",
            "机体可再用（秒）",
            "电池可再用（秒）",
            "能耗（kWh）",
            "返航电量（%）",
            "箱数",
            "货箱编号列表",
        ],
        [
            [
                i + 1,
                r["sortie_id"],
                r["drone_id"],
                drones.get(r["drone_id"]),
                r["battery_id"],
                r["prepare_start_s"],
                r["takeoff_s"],
                " → ".join(label(v["site_id"]) for v in r["visits"]),
                r["return_s"],
                r["drone_released_s"],
                r["battery_released_s"],
                r["energy_kwh"],
                100 * r["return_soc"],
                len(r["cargo_ids"]),
                "; ".join(r["cargo_ids"]),
            ]
            for i, r in enumerate(ordered)
        ],
    )
    deliveries = []
    for r in ordered:
        for visit in r["visits"]:
            for cid in visit["cargo_ids"]:
                box = cargo[cid]
                d = demands[box["demand_id"]]
                expected = d.get("expected_deadline_s")
                hard = d.get("hard_deadline_s")
                deliveries.append(
                    [
                        cid,
                        r["sortie_id"],
                        label(visit["site_id"]),
                        r["drone_id"],
                        r["battery_id"],
                        types.get(box["type"], box["type"]),
                        box["mass_kg"],
                        visit["delivery_complete_s"],
                        expected,
                        hard,
                        (
                            max(0, visit["delivery_complete_s"] - expected)
                            if expected is not None
                            else None
                        ),
                    ]
                )
    sheet(
        "逐箱交付",
        [
            "货箱编号",
            "架次编号",
            "救援点",
            "无人机编号",
            "电池编号",
            "物资类别",
            "质量（kg）",
            "交付完成（秒）",
            "期望送达（秒）",
            "强制截止（秒）",
            "超过期望（秒）",
        ],
        sorted(deliveries, key=lambda x: x[0]),
    )
    relay = plan.get("relay_sorties", [])
    sheet(
        "中继架次",
        [
            "中继架次编号",
            "中继无人机编号",
            "能源组件编号",
            "准备开始（秒）",
            "悬停经度",
            "悬停纬度",
            "悬停海拔（m）",
            "离地高度（m）",
            "建链完成（秒）",
            "服务结束（秒）",
            "返回调度中心（秒）",
            "能耗（kWh）",
            "返航电量（%）",
        ],
        [
            [
                r["relay_sortie_id"],
                r["drone_id"],
                r["component_id"],
                r["prepare_start_s"],
                *r["station_coordinates"],
                r["station_elevation_m"],
                r["station_agl_m"],
                r["service_start_s"],
                r["service_end_s"],
                r["return_s"],
                r["energy_kwh"],
                100 * r["return_soc"],
            ]
            for r in relay
        ],
    )
    windows = (plan.get("communication_report") or {}).get("windows", [])
    phase_names = {
        "climb": "爬升",
        "cruise": "巡航",
        "descent": "下降",
        "handover": "物资交接",
    }
    sheet(
        "通信保障",
        [
            "运输架次编号",
            "通信阶段",
            "开始时刻（秒）",
            "结束时刻（秒）",
            "保障方式",
            "中继架次编号",
        ],
        [
            [
                w["sortie_id"],
                " / ".join(
                    dict.fromkeys(
                        phase_names.get(s["phase"], s["phase"])
                        for s in phases
                        if s["sortie_id"] == w["sortie_id"]
                        and s["start_s"] < w["end_s"]
                        and s["end_s"] > w["start_s"]
                    )
                ),
                w["start_s"],
                w["end_s"],
                {"gateway": "直连地面网关", "relay": "中继转发"}.get(
                    w["source"], w["source"]
                ),
                w.get("relay_sortie_id"),
            ]
            for w in windows
        ],
    )
    resources = []
    for r in ordered:
        resources.extend(
            [
                [
                    "运输无人机",
                    r["drone_id"],
                    r["sortie_id"],
                    r["prepare_start_s"],
                    r["drone_released_s"],
                ],
                [
                    "运输电池",
                    r["battery_id"],
                    r["sortie_id"],
                    r["prepare_start_s"],
                    r["battery_released_s"],
                ],
            ]
        )
    for r in relay:
        resources.extend(
            [
                [
                    "中继无人机",
                    r["drone_id"],
                    r["relay_sortie_id"],
                    r["prepare_start_s"],
                    r["drone_released_s"],
                ],
                [
                    "中继能源组件",
                    r["component_id"],
                    r["relay_sortie_id"],
                    r["prepare_start_s"],
                    r["component_released_s"],
                ],
            ]
        )
    sheet(
        "资源占用",
        ["资源类型", "实体编号", "架次编号", "占用开始（秒）", "再次可用（秒）"],
        resources,
    )
    output = BytesIO()
    wb.save(output)
    return output.getvalue()
