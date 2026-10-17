"""Bounded JSON/CSV/XLSX scene import and editable template export."""

import copy
import csv
import io
import json
import zipfile

from openpyxl import Workbook, load_workbook
from pydantic import ValidationError

from backend.app.generated.domain import Scenario
from shared.contracts import validate_scenario

GROUPS = ["sites", "demands", "cargo_items", "drone_types", "drones", "batteries"]
OBJECTS = {"terrain", "base", "policy", "energy_model", "charge_model", "coordinates"}
NUMBERS = {
    "state_version",
    "ground_elevation_m",
    "required_quantity",
    "priority",
    "expected_deadline_s",
    "hard_deadline_s",
    "mass_kg",
    "volume_l",
    "available_after_s",
    "capacity_kwh",
    "soc",
    "soh",
    "max_payload_kg",
    "max_volume_l",
    "cruise_speed_mps",
    "climb_speed_mps",
    "descent_speed_mps",
    "prepare_duration_s",
    "turnaround_duration_s",
    "handover_fixed_s",
    "handover_per_box_s",
    "empty_mass_kg",
}
INTS = {"state_version", "required_quantity", "priority"}
BOOLS = {"available", "mandatory"}


def parse_value(key, value):
    if value is None or value == "":
        return None
    if key in OBJECTS:
        return json.loads(value) if isinstance(value, str) else value
    if key in BOOLS:
        if isinstance(value, bool):
            return value
        text = str(value).lower().strip()
        if text not in ("true", "false", "1", "0", "是", "否"):
            raise ValueError(f"{key}需要true/false")
        return text in ("true", "1", "是")
    if key in NUMBERS:
        if key in INTS:
            n = float(value)
            if not n.is_integer():
                raise ValueError(f"{key}必须为整数")
            return int(n)
        return float(value)
    return str(value)


def parse_rows(rows):
    if not rows:
        return []
    header = [str(v).strip() if v is not None else "" for v in rows[0]]
    if len(header) != len(set(header)) or any(not v for v in header):
        raise ValueError("表头存在空字段或重复字段。")
    result = []
    for i, row in enumerate(rows[1:], start=2):
        if not any(v not in ("", None) for v in row):
            continue
        try:
            result.append({k: parse_value(k, v) for k, v in zip(header, row)})
        except (ValueError, TypeError) as exc:
            raise ValueError(f"第{i}行：{exc}") from exc
    return result


def preview(filename, data, current=None, entity=None):
    suffix = filename.lower().rsplit(".", 1)[-1]
    if suffix == "json":
        value = json.loads(data.decode("utf-8-sig"))
        s = value.get("scenario", value) if isinstance(value, dict) else value
    elif suffix == "csv":
        if current is None or entity not in GROUPS:
            raise ValueError("CSV导入须选择当前场景和实体表。")
        s = copy.deepcopy(current)
        rows = list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))
        s[entity] = parse_rows(rows)
    elif suffix == "xlsx":
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if (
                len(entries) > 2000
                or sum(e.file_size for e in entries) > 50 * 1024 * 1024
            ):
                raise ValueError("工作簿解压规模超出限制。")
            if any(".." in e.filename.split("/") for e in entries):
                raise ValueError("工作簿包含不安全路径。")
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            if not all(name in wb.sheetnames for name in ["scenario", *GROUPS]):
                raise ValueError("请使用援航完整场景Excel模板。")
            metadata = parse_rows([list(r) for r in wb["scenario"].values])
            if len(metadata) != 1:
                raise ValueError("scenario工作表必须只有一条场景记录。")
            s = metadata[0]
            for group in GROUPS:
                s[group] = parse_rows([list(r) for r in wb[group].values])
        finally:
            wb.close()
    else:
        raise ValueError("仅支持JSON、UTF-8 CSV和XLSX。")
    if not isinstance(s, dict):
        raise ValueError("场景必须为JSON对象。")
    if s.get("data_origin") != "mock":
        s["data_origin"] = "user_input"
    normalized = Scenario.model_validate(s).model_dump(mode="json")
    validate_scenario(normalized)
    return {
        "scenario": normalized,
        "row_count": sum(len(normalized[g]) for g in GROUPS),
        "issues": [],
    }


def csv_template(scenario, entity):
    rows = scenario[entity]
    buffer = io.StringIO()
    keys = list(rows[0]) if rows else []
    writer = csv.writer(buffer)
    writer.writerow(keys)
    for row in rows:
        writer.writerow(
            [
                (
                    json.dumps(row[k], ensure_ascii=False)
                    if isinstance(row[k], (list, dict))
                    else row[k]
                )
                for k in keys
            ]
        )
    return ("\ufeff" + buffer.getvalue()).encode("utf-8")


def excel_template(scenario):
    wb = Workbook()
    wb.remove(wb.active)
    for group, rows in [
        ("scenario", [{k: v for k, v in scenario.items() if k not in GROUPS}]),
        *[(g, scenario[g]) for g in GROUPS],
    ]:
        ws = wb.create_sheet(group)
        if not rows:
            continue
        keys = list(rows[0])
        ws.append(keys)
        for row in rows:
            ws.append(
                [
                    (
                        json.dumps(row[k], ensure_ascii=False)
                        if isinstance(row[k], (list, dict))
                        else row[k]
                    )
                    for k in keys
                ]
            )
        ws.freeze_panes = "A2"
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = 22
    buffer = io.BytesIO()
    wb.save(buffer)
    wb.close()
    return buffer.getvalue()
